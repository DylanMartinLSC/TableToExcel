"""Local extraction and safe workbook export; no network services."""

import csv
import os
import re
import shutil
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageOps

MAX_BYTES = 30 * 1024 * 1024
MAX_PIXELS = 40_000_000
MAX_PAGES = 50
EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".pdf"}


@dataclass
class Table:
    name: str
    rows: list[list[str]]
    page: int = 1
    merges: list[tuple[int, int, int, int]] = field(default_factory=list)


def setup_tesseract():
    location = Path(r"C:\Program Files\Tesseract-OCR")
    if not shutil.which("tesseract") and (location / "tesseract.exe").exists():
        os.environ["PATH"] = str(location) + os.pathsep + os.environ.get("PATH", "")
    if not shutil.which("tesseract"):
        raise ValueError("Install Tesseract OCR, then restart this app. See the installation guide.")


def parse_pages(value: str) -> list[int] | None:
    if not value.strip():
        return None
    pages = set()
    try:
        for part in value.split(","):
            ends = [int(x.strip()) for x in part.split("-")]
            if len(ends) not in (1, 2) or min(ends) < 1 or max(ends) > MAX_PAGES:
                raise ValueError
            first, last = ends[0], ends[-1]
            if first > last:
                raise ValueError
            pages.update(range(first - 1, last))
    except ValueError:
        raise ValueError("Use page numbers such as 1, 3-5 (maximum 50).") from None
    return sorted(pages)


def prepare(data: bytes, name: str, rotation=0, crop=None):
    ext = Path(name).suffix.lower()
    if ext not in EXTENSIONS:
        raise ValueError("Choose a PNG, JPEG, WebP, BMP, TIFF or PDF file.")
    if not data or len(data) > MAX_BYTES:
        raise ValueError("Each file must be nonempty and no larger than 30 MB.")
    if ext == ".pdf":
        import pypdfium2 as pdfium

        try:
            with pdfium.PdfDocument(data) as pdf:
                count = len(pdf)
        except Exception:
            raise ValueError("Cannot open this PDF. It may be damaged or password-protected.") from None
        if not 0 < count <= MAX_PAGES:
            raise ValueError("PDFs must have between 1 and 50 pages. Split larger files first.")
        return data, count
    try:
        with Image.open(BytesIO(data)) as original:
            if getattr(original, "n_frames", 1) > 1:
                raise ValueError("Multi-frame images: save each page separately or use PDF.")
            if original.width * original.height > MAX_PIXELS:
                raise ValueError("Image is too large. Resize below 40 megapixels.")
            image = ImageOps.exif_transpose(original).convert("RGB")
    except (OSError, Image.DecompressionBombError):
        raise ValueError("Cannot read this image. Try saving it as PNG.") from None
    image = image.rotate(rotation, expand=True)
    if crop:
        left, top, right, bottom = crop
        if not (0 <= left < right <= 100 and 0 <= top < bottom <= 100):
            raise ValueError("Crop edges must enclose a nonempty area within 0–100%.")
        image = image.crop(
            (
                round(left * image.width / 100),
                round(top * image.height / 100),
                round(right * image.width / 100),
                round(bottom * image.height / 100),
            )
        )
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue(), 1


def convert_table(raw, name, page):
    cells = list(raw.content.values())
    rows = [[str(c.value) if c.value is not None else "" for c in row] for row in cells]
    groups = {}
    for r, row in enumerate(cells, 1):
        for c, cell in enumerate(row, 1):
            b = cell.bbox
            groups.setdefault((b.x1, b.y1, b.x2, b.y2), []).append((r, c))
    merges = []
    for positions in groups.values():
        if len(positions) > 1:
            rs, cs = zip(*positions)
            if len(positions) == (max(rs) - min(rs) + 1) * (max(cs) - min(cs) + 1):
                merges.append((min(rs), min(cs), max(rs), max(cs)))
    return Table(name, rows, page, merges)


def remove_rules(image):
    """Blank long table rules so OCR sees only text; the original keeps the structure."""
    import cv2

    gray = cv2.cvtColor(image, cv2.COLOR_RGB2GRAY)
    ink = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)[1]
    horizontal = cv2.morphologyEx(
        ink,
        cv2.MORPH_OPEN,
        cv2.getStructuringElement(cv2.MORPH_RECT, (max(40, image.shape[1] // 20), 1)),
    )
    vertical = cv2.morphologyEx(
        ink,
        cv2.MORPH_OPEN,
        cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(40, image.shape[0] // 15))),
    )
    mask = cv2.dilate(horizontal | vertical, cv2.getStructuringElement(cv2.MORPH_RECT, (3, 3)))
    cleaned = image.copy()
    cleaned[mask > 0] = 255
    return cleaned


def extract(
    data,
    name,
    *,
    pages=None,
    language="eng",
    borderless=False,
    confidence=50,
    implicit=False,
    rotation=0,
    crop=None,
):
    from img2table.document import PDF
    from img2table.document import Image as TableImage
    from img2table.ocr import TesseractOCR

    setup_tesseract()
    data, count = prepare(data, name, rotation, crop)
    if pages and max(pages) >= count:
        raise ValueError(f"This document has only {count} page(s).")
    if not re.fullmatch(r"[A-Za-z0-9_]+(?:\+[A-Za-z0-9_]+)*", language):
        raise ValueError("Use language codes such as eng or eng+fra.")

    class LocalOCR(TesseractOCR):
        def hocr(self, image):
            import subprocess
            import tempfile

            with tempfile.TemporaryDirectory(prefix="tabletoexcel-") as folder:
                path = Path(folder) / "ocr.png"
                Image.fromarray(remove_rules(image)).save(path)
                return subprocess.check_output(
                    ["tesseract", str(path), "stdout", "--psm", "11", "-l", self.lang, "hocr"],
                    env=self.env,
                    stderr=subprocess.PIPE,
                    timeout=120,
                ).decode("utf-8")

    ocr = LocalOCR(n_threads=1, lang=language)
    pdf = Path(name).suffix.lower() == ".pdf"
    doc = PDF(data, pages=pages, detect_rotation=True) if pdf else TableImage(data, detect_rotation=True)
    result = doc.extract_tables(
        ocr=ocr,
        implicit_rows=implicit,
        implicit_columns=implicit,
        borderless_tables=borderless,
        min_confidence=confidence,
    )
    grouped = result if pdf else {0: result}
    return [
        convert_table(t, f"{Path(name).stem} p{p + 1} t{i + 1}", p + 1)
        for p, tables in grouped.items()
        for i, t in enumerate(tables)
    ]


def workbook(tables: list[Table], preserve_merges=True) -> bytes:
    if not tables:
        raise ValueError("No tables to export.")
    book = Workbook()
    book.remove(book.active)
    used = set()
    for table in tables:
        base = re.sub(r"[\\/*?:\[\]]", "_", table.name).strip("'")[:31] or "Table"
        name, suffix = base, 1
        while name.lower() in used:
            suffix += 1
            name = base[: 31 - len(str(suffix)) - 1] + "_" + str(suffix)
        used.add(name.lower())
        sheet = book.create_sheet(name)
        for row in table.rows:
            sheet.append([str(v) if v is not None else "" for v in row])
        for row in sheet:
            for cell in row:
                cell.data_type = "s"  # OCR/user text must never become executable formulas.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="174B56")
        if preserve_merges:
            for r1, c1, r2, c2 in table.merges:
                if r2 <= sheet.max_row and c2 <= sheet.max_column:
                    sheet.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        from openpyxl.utils import get_column_letter

        for col in range(1, sheet.max_column + 1):
            width = max(
                (len(str(sheet.cell(r, col).value or "")) for r in range(1, sheet.max_row + 1)), default=10
            )
            sheet.column_dimensions[get_column_letter(col)].width = min(55, max(12, width + 2))
        sheet.freeze_panes = "A2"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def csv_bytes(table):
    from io import StringIO

    output = StringIO(newline="")
    writer = csv.writer(output)
    for row in table.rows:
        writer.writerow(
            ["'" + str(v) if str(v).lstrip().startswith(("=", "+", "-", "@")) else v for v in row]
        )
    return output.getvalue().encode("utf-8-sig")
