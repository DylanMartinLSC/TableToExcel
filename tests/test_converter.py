from io import BytesIO
from types import SimpleNamespace as NS

import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

from converter import Table, convert_table, csv_bytes, extract, parse_pages, prepare, workbook


@pytest.mark.parametrize("text,result", [("", None), ("1, 3-5", [0, 2, 3, 4]), ("2,2", [1])])
def test_pages(text, result):
    assert parse_pages(text) == result


@pytest.mark.parametrize("text", ["0", "51", "3-1", "x", "1-2-3", ","])
def test_bad_pages(text):
    with pytest.raises(ValueError):
        parse_pages(text)


@pytest.mark.parametrize(
    "name,data", [("x.exe", b"abc"), ("x.png", b"bad"), ("x.pdf", b"bad"), ("x.png", b"")]
)
def test_bad_inputs(name, data):
    with pytest.raises(ValueError):
        prepare(data, name)


def test_crop_rotation():
    data = BytesIO()
    Image.new("RGB", (100, 200)).save(data, format="PNG")
    prepared, _ = prepare(data.getvalue(), "x.png", 90, (0, 0, 50, 100))
    assert Image.open(BytesIO(prepared)).size == (100, 100)
    with pytest.raises(ValueError):
        prepare(data.getvalue(), "x.png", crop=(90, 0, 10, 100))


def test_export():
    tables = [
        Table("bad/name", [["=1+1", "0012"], ["a", "b"]], merges=[(2, 1, 2, 2)]),
        Table("bad/name", [["x"]]),
    ]
    book = load_workbook(BytesIO(workbook(tables)))
    assert book.sheetnames == ["bad_name", "bad_name_2"]
    assert book.worksheets[0]["A1"].data_type == "s"
    assert book.worksheets[0]["B1"].value == "0012"
    assert str(book.worksheets[0].merged_cells) == "A2:B2"
    assert not load_workbook(BytesIO(workbook(tables, False))).worksheets[0].merged_cells
    assert b"'=1+1" in csv_bytes(tables[0])


def test_empty_export():
    with pytest.raises(ValueError):
        workbook([])


def test_merges():
    a = NS(value="Title", bbox=NS(x1=0, y1=0, x2=100, y2=10))
    raw = NS(content={0: [a, a]})
    assert convert_table(raw, "test", 1).merges == [(1, 1, 1, 2)]


@pytest.mark.integration
@pytest.mark.parametrize("kind", ["png", "pdf"])
def test_real_ocr(tmp_path, kind):
    image = Image.new("RGB", (1200, 440), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 32)
    except OSError:
        try:
            font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 32)
        except OSError:
            font = ImageFont.load_default(size=32)
    rows = [["Product", "Quantity", "Price"], ["Apples", "12", "24"], ["Oranges", "8", "16"]]
    for x in [30, 410, 790, 1170]:
        draw.line((x, 30, x, 390), fill="black", width=3)
    for y in [30, 150, 270, 390]:
        draw.line((30, y, 1170, y), fill="black", width=3)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            draw.text((50 + c * 380, 65 + r * 120), value, font=font, fill="black")
    source = tmp_path / f"sample.{kind}"
    image.save(source, resolution=200)
    result = extract(source.read_bytes(), source.name)
    assert len(result) == 1
    text = " ".join(v for row in result[0].rows for v in row)
    assert "Apples" in text and "Oranges" in text and "24" in text
    assert len(result[0].rows) == 3
    (tmp_path / "sample.xlsx").write_bytes(workbook(result))


def test_app_review():
    from streamlit.testing.v1 import AppTest

    app = AppTest.from_file(str(__import__("pathlib").Path("app.py").resolve())).run()
    assert not app.exception
    app.session_state["tables"] = [Table("Example", [["Name", "Count"], ["Apples", "12"]])]
    app.run()
    assert not app.exception
    assert len(app.get("download_button")) == 2
